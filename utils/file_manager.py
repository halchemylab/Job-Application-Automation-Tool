import os
import re
from pathlib import Path
from docx import Document

# Base output directory
OUTPUT_DIR = Path("output")

# Ensure the output directory exists
OUTPUT_DIR.mkdir(exist_ok=True)

def sanitize_name(name):
    """
    Sanitizes a string to make it suitable for file and folder names.
    """
    return re.sub(r"[^A-Za-z0-9_-]+", "_", name.strip())

def create_application_folder(company_name, job_title):
    """
    Creates a folder for a specific job application.

    :param company_name: The name of the company
    :param job_title: The job title
    :return: Path to the created folder
    """
    folder_name = f"{sanitize_name(company_name)}-{sanitize_name(job_title)}"
    folder_path = OUTPUT_DIR / folder_name
    folder_path.mkdir(parents=True, exist_ok=True)
    return folder_path

def save_customized_document(template_path, output_path, replacements):
    """
    Customizes a Word document based on a template and saves it.

    :param template_path: Path to the template Word document
    :param output_path: Path to save the customized document
    :param replacements: A dictionary of placeholders and their replacements
    """
    # Load the template
    doc = Document(template_path)

    # Replace placeholders in the document
    for paragraph in doc.paragraphs:
        for placeholder, replacement in replacements.items():
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, replacement)

    # Save the customized document
    doc.save(output_path)

def save_application_files(company_name, job_title, job_description, replacements):
    """
    Creates the application folder and saves the customized resume and cover letter.

    :param company_name: The name of the company
    :param job_title: The job title
    :param job_description: The job description text (not currently used here but available for future use)
    :param replacements: A dictionary of placeholders to be replaced in templates
    :return: Path to the application folder
    """
    # Create the application folder
    application_folder = create_application_folder(company_name, job_title)

    # Define paths to templates
    resume_template_path = Path("static/resume_template.docx")
    cover_letter_template_path = Path("static/cover_letter_template.docx")

    # Define output paths
    resume_output_path = application_folder / "resume.docx"
    cover_letter_output_path = application_folder / "cover_letter.docx"

    # Save customized documents
    save_customized_document(resume_template_path, resume_output_path, replacements)
    save_customized_document(cover_letter_template_path, cover_letter_output_path, replacements)

    return application_folder

# Example usage
if __name__ == "__main__":
    # Example data
    company_name = "Awesome Tech"
    job_title = "Data Scientist"
    job_description = "Analyze and interpret complex data to help decision-making."

    # Example replacements for templates
    replacements = {
        "[COMPANY_NAME]": company_name,
        "[JOB_TITLE]": job_title,
        "[JOB_DESCRIPTION]": job_description
    }

    folder = save_application_files(company_name, job_title, job_description, replacements)
    print(f"Files saved in: {folder}")
